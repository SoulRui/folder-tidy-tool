# -*- coding: utf-8 -*-
"""
文件夹整理工具
==============
功能一：按 Excel 分类整理
  - 整理对象二选一：对文件夹调整（移动子文件夹）/ 对里面的文件调整（移动文件，不限格式）
  - 处理方式二选一：移动（在待处理文件夹内整理）/ 复制到新目录（原文件保留）
  - 自动识别表头行与列名（"编号/文件夹名/名称/ID…"、"一级/二级/三级…"等相近写法均可）
  - 名称中**包含**编号字符串即匹配：编号可为文字+数字等任意字符串、位数不限；
    位置不限（开头/中间/结尾均可）；同时包含多个编号时取最长者
  - 去向按 Excel 级别目录（一级/二级/三级/…，级别数不限，取最深非空级）
  - 未匹配到的源文件夹/文件进入"整理失败"文件夹，并生成《整理结果报告》Excel

功能二：批量获取文件夹名（生成目录清单）
  - 递归扫描所选文件夹，列出全部子文件夹与文件（含分级路径列）
  - 可选：同时与分类 Excel 比对，标注每个顶层文件夹的归类去向

功能二·五：按 Excel 批量重命名
  - Excel 两列：一列旧名称、一列新名称（自动识别表头；无表头按 A 列旧、B 列新）
  - 整理对象二选一：只重命名子文件夹 / 只重命名文件（不限格式，按去后缀名匹配，
    新名称不带扩展名时保留原扩展名）
  - 处理方式二选一：移动（原位重命名）/ 复制到新目录（原位置保留）
  - 生成《重命名结果报告》Excel

功能三：批量移动同名文件到同一文件夹
  - 按完整文件名分组：name.shp/name.shx/name.dbf … 归入 name 文件夹
  - 按文件名前 N 位分组：如取前 5 位字符作为文件夹名
  - 可限定扩展名（留空处理全部文件）

支持拖拽导入文件/文件夹、自动记住上次路径。
"""
import os
import re
import sys
import json
import shutil
import threading
import traceback
from datetime import datetime

APP_NAME = "文件夹整理工具"
FAIL_DIR_NAME = "整理失败"


def _resource_path(relative_path):
    base = getattr(sys, "_MEIPASS", os.path.dirname(os.path.abspath(__file__)))
    return os.path.join(base, relative_path)


# ---------------------------------------------------------------- 配置记忆

def _config_path():
    base = os.environ.get("APPDATA") or os.path.expanduser("~")
    d = os.path.join(base, APP_NAME)
    os.makedirs(d, exist_ok=True)
    return os.path.join(d, "config.json")


def load_config():
    try:
        with open(_config_path(), "r", encoding="utf-8") as f:
            return json.load(f)
    except Exception:
        return {}


def save_config(cfg):
    try:
        with open(_config_path(), "w", encoding="utf-8") as f:
            json.dump(cfg, f, ensure_ascii=False, indent=2)
    except Exception:
        pass


# ---------------------------------------------------------------- Excel 解析

CN_NUM = {"一": 1, "二": 2, "三": 3, "四": 4, "五": 5,
          "六": 6, "七": 7, "八": 8, "九": 9, "十": 10}

# 编号列的相近叫法（归一化后包含任一关键词即认为是对应列）
ID_KEYWORDS = ["编号", "文件夹名", "文件名", "名称", "原名", "图版", "图幅",
               "图号", "目录名", "id", "code", "name"]

LEVEL_RE = re.compile(r"^第?([一二三四五六七八九十]|\d+)\s*级")


def _norm(s):
    return re.sub(r"[\s_（）()\[\]【】]", "", str(s)).lower() if s is not None else ""


def _match_id_col(cell_norm):
    return any(k in cell_norm for k in ID_KEYWORDS)


def _match_level_col(cell_norm):
    m = LEVEL_RE.match(cell_norm)
    if not m:
        return None
    token = m.group(1)
    return CN_NUM.get(token, int(token) if token.isdigit() else None)


def find_header(ws, scan_rows=10):
    """在前 scan_rows 行中定位表头行。
    返回 (表头行号[1-based], 编号列索引[0-based], [(层级数, 列索引), ...])"""
    best = None
    for r_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=scan_rows, values_only=True), start=1):
        id_col = None
        levels = []
        for c_idx, v in enumerate(row):
            cn = _norm(v)
            if not cn:
                continue
            if id_col is None and _match_id_col(cn):
                id_col = c_idx
                continue
            lv = _match_level_col(cn)
            if lv is not None:
                levels.append((lv, c_idx))
        if id_col is not None and levels:
            levels.sort()
            score = len(levels)
            if best is None or score > best[0]:
                best = (score, r_idx, id_col, levels)
    if best is None:
        raise ValueError("未能在前 %d 行识别到表头：需要包含『编号/文件夹名/名称』类列和『一级/二级…』类列" % scan_rows)
    _, r_idx, id_col, levels = best
    return r_idx, id_col, [(lv, ci) for lv, ci in levels]


def _cell_to_name(v):
    """单元格 -> 文件夹名字符串（数字编号去小数点，文字原样保留）"""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def load_classification(excel_path, log=print):
    """读取分类 Excel，返回 {原文件夹名: [一级, 二级, ...]}（空层级为 None，按层级排序）"""
    import openpyxl
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    last_err = None
    for ws in wb.worksheets:  # 逐个 Sheet 尝试，取第一个能识别表头的
        try:
            header_row, id_col, level_cols = find_header(ws)
        except ValueError as e:
            last_err = e
            continue
        header = [ws.cell(row=header_row, column=c + 1).value for c in range(ws.max_column)]
        log(f"使用 Sheet「{ws.title}」，表头在第 {header_row} 行")
        log(f"识别列：编号列=「{header[id_col]}」，层级列="
            + "、".join(f"「{header[ci]}」(第{lv}级)" for lv, ci in level_cols))
        mapping = {}
        for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
            name = _cell_to_name(row[id_col] if id_col < len(row) else None)
            if not name:
                continue
            levels = []
            for _, ci in level_cols:
                lv = _cell_to_name(row[ci] if ci < len(row) else None)
                levels.append(lv if lv else None)
            mapping[name] = levels
        return mapping
    raise ValueError(str(last_err) if last_err else "Excel 中没有可用的 Sheet")


def deepest_path(levels):
    parts = [lv for lv in levels if lv]
    return os.path.join(*parts) if parts else None


def match_key(name, mapping):
    """在分类表编号中查找与 name 匹配的编号。
    优先完全相等；否则取被 name **包含**的最长编号（编号可在开头/中间/结尾）。
    匹配不区分大小写。返回分类表中的原始编号，未匹配返回 None。"""
    if not name:
        return None
    if name in mapping:
        return name
    ln = name.lower()
    best = None
    for k in mapping:
        kl = k.lower()
        if not kl:
            continue
        if kl == ln:
            return k
        if kl in ln and (best is None or len(kl) > len(best)):
            best = kl
    if best is None:
        return None
    for k in mapping:  # 还原原始大小写
        if k.lower() == best:
            return k
    return None


# ---------------------------------------------------------------- 功能一：分类整理

def classify(excel_path, work_dir, mode="both", action="move", out_dir=None, log=print):
    """mode: 'dirs' 只整理文件夹；'files' 只整理文件；'both' 两者都整理
    action: 'move' 移动（在待处理文件夹内整理）；'copy' 复制到新目录（需 out_dir）"""
    import openpyxl
    if action == "copy":
        if not out_dir:
            raise ValueError("复制模式需要指定输出目录")
        os.makedirs(out_dir, exist_ok=True)
        base_dir = out_dir          # 目标根：新目录
    else:
        action = "move"
        base_dir = work_dir         # 目标根：待处理文件夹本身

    mapping = load_classification(excel_path, log=log)
    log(f"分类表共 {len(mapping)} 条记录")

    # 建立所需目录
    needed = {deepest_path(lv) for lv in mapping.values()}
    needed.discard(None)
    for rel in sorted(needed):
        os.makedirs(os.path.join(base_dir, rel), exist_ok=True)
    log(f"已在 {base_dir} 建立 {len(needed)} 个目标目录")

    category_names = set()
    for lv in mapping.values():
        category_names.update(x for x in lv if x)

    # 扫描待处理文件夹顶层（而不是仅按 Excel 行），保证未匹配的也能被发现
    # 文件夹按名称匹配；文件不限格式，名称中包含编号字符串即匹配（位置不限）
    SKIP_FILE_PREFIX = ("整理结果报告", "目录清单")
    SKIP_FILE_NAMES = {"desktop.ini", "thumbs.db"}
    top_dirs, top_files = [], []
    for d in sorted(os.listdir(work_dir)):
        p = os.path.join(work_dir, d)
        if os.path.isdir(p):
            if d != FAIL_DIR_NAME and d not in category_names:
                top_dirs.append(d)
        elif os.path.isfile(p):
            if d.startswith(SKIP_FILE_PREFIX) or d.lower() in SKIP_FILE_NAMES:
                continue
            top_files.append(d)
    do_dirs = mode in ("both", "dirs")
    do_files = mode in ("both", "files")
    mode_text = {"both": "文件夹+文件", "dirs": "仅文件夹", "files": "仅文件"}.get(mode, mode)
    action_text = "复制到新目录" if action == "copy" else "移动（在待处理文件夹内整理）"
    log(f"整理对象：{mode_text}；处理方式：{action_text}")
    log(f"待处理顶层文件夹 {len(top_dirs)} 个，文件 {len(top_files)} 个")

    records = []  # (名称, 处理结果, 目标位置, 说明)
    moved = failed = conflict = 0              # 文件夹
    moved_f = failed_f = conflict_f = 0        # 文件
    matched_names = set()
    fail_dir = os.path.join(base_dir, FAIL_DIR_NAME)
    verb = "复制" if action == "copy" else "移动"

    def _transfer(src, dst, is_dir):
        if action == "copy":
            if is_dir:
                shutil.copytree(src, dst)
            else:
                shutil.copy2(src, dst)
        else:
            shutil.move(src, dst)

    for name in top_dirs if do_dirs else []:
        key = match_key(name, mapping)
        src = os.path.join(work_dir, name)
        if key is None:
            # 未匹配 -> 整理失败
            os.makedirs(fail_dir, exist_ok=True)
            dst = os.path.join(fail_dir, name)
            if os.path.exists(dst):
                records.append((name, f"未匹配-未{verb}", FAIL_DIR_NAME, "整理失败中已存在同名文件夹"))
                conflict += 1
                log(f"[冲突] {name}：整理失败文件夹中已存在同名项，未{verb}")
                continue
            _transfer(src, dst, is_dir=True)
            failed += 1
            records.append((name, f"未匹配-已{verb}到整理失败", FAIL_DIR_NAME, "分类表中没有编号被该名称包含"))
            log(f"[未匹配] {name} -> {FAIL_DIR_NAME}")
            continue
        matched_names.add(key)
        note = "" if key == name else f"按包含的编号「{key}」匹配"
        rel = deepest_path(mapping[key])
        if not rel:
            records.append((name, f"未{verb}", "", "该条记录层级均为空"))
            log(f"[跳过] {name}：分类表中层级均为空")
            continue
        dst = os.path.join(base_dir, rel, name)
        if os.path.exists(dst):
            conflict += 1
            records.append((name, f"目标冲突-未{verb}", rel, "目标位置已存在同名文件夹"))
            log(f"[冲突] {name}：目标 {rel} 已存在同名文件夹")
            continue
        _transfer(src, dst, is_dir=True)
        moved += 1
        records.append((name, f"已{verb}", rel, note))
        log(f"[{verb}] {name} -> {rel}" + (f"（{note}）" if note else ""))

    # ---------------- 顶层文件：名称中包含编号字符串即匹配（不限格式、编号为任意字符串） ----------------
    for fn in top_files if do_files else []:
        stem = os.path.splitext(fn)[0].strip()
        key = match_key(stem, mapping) or match_key(fn, mapping)  # 先按去后缀名，再按完整文件名
        src = os.path.join(work_dir, fn)
        if key is None:
            os.makedirs(fail_dir, exist_ok=True)
            dst = os.path.join(fail_dir, fn)
            if os.path.exists(dst):
                records.append((fn, f"未匹配-未{verb}", FAIL_DIR_NAME, "整理失败中已存在同名文件"))
                conflict_f += 1
                log(f"[冲突] {fn}：整理失败文件夹中已存在同名文件，未{verb}")
                continue
            _transfer(src, dst, is_dir=False)
            failed_f += 1
            records.append((fn, f"未匹配-已{verb}到整理失败", FAIL_DIR_NAME,
                            "分类表中没有编号被该文件名包含"))
            log(f"[未匹配] {fn} -> {FAIL_DIR_NAME}")
            continue
        matched_names.add(key)
        note = "" if key == stem else f"按包含的编号「{key}」匹配"
        rel = deepest_path(mapping[key])
        if not rel:
            records.append((fn, f"未{verb}", "", "该条记录层级均为空"))
            log(f"[跳过] {fn}：分类表中层级均为空")
            continue
        # 直接归入表格中该编号对应的最深层级目录（一级/二级/三级/…），不建编号子文件夹
        dst_dir = os.path.join(base_dir, rel)
        dst = os.path.join(dst_dir, fn)
        if os.path.exists(dst):
            conflict_f += 1
            records.append((fn, f"目标冲突-未{verb}", rel, "目标位置已存在同名文件"))
            log(f"[冲突] {fn}：目标 {rel} 已存在同名文件")
            continue
        os.makedirs(dst_dir, exist_ok=True)
        _transfer(src, dst, is_dir=False)
        moved_f += 1
        records.append((fn, f"已{verb}", rel, note))
        log(f"[{verb}] {fn} -> {rel}" + (f"（{note}）" if note else ""))

    missing = 0
    for key in sorted(set(mapping) - matched_names):
        if not os.path.isdir(os.path.join(work_dir, key)) and key not in category_names:
            # 可能在上一轮已移动；仅当分类目录里也没有时记为缺失
            rel = deepest_path(mapping[key])
            if rel and os.path.isdir(os.path.join(base_dir, rel, key)):
                continue
            missing += 1
            records.append((key, "源文件夹缺失", deepest_path(mapping[key]) or "", "分类表有记录，但待处理文件夹中不存在"))
    if missing:
        log(f"分类表中有 {missing} 条记录在待处理文件夹中找不到对应文件夹（详见报告）")

    # 结果报告
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    report_path = os.path.join(base_dir, f"整理结果报告_{ts}.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "整理结果"
    ws.append(["原文件夹/文件名", "处理结果", "目标位置", "说明"])
    for rec in records:
        ws.append(list(rec))
    summary = wb.create_sheet("汇总")
    summary.append(["项目", "数量"])
    summary.append(["整理对象", mode_text])
    summary.append(["处理方式", action_text])
    summary.append([f"成功{verb}文件夹", moved])
    summary.append([f"成功{verb}文件", moved_f])
    summary.append([f"未匹配文件夹（{verb}到整理失败）", failed])
    summary.append([f"未匹配文件（{verb}到整理失败）", failed_f])
    summary.append([f"冲突未{verb}（文件夹/文件）", f"{conflict} / {conflict_f}"])
    summary.append(["分类表有记录但源缺失", missing])
    summary.append(["待处理顶层文件夹数", len(top_dirs)])
    summary.append(["待处理顶层文件数", len(top_files)])
    summary.append(["报告生成时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")])
    for sheet in (ws, summary):
        for col_cells in sheet.columns:
            width = max(len(str(c.value)) if c.value is not None else 0 for c in col_cells)
            sheet.column_dimensions[col_cells[0].column_letter].width = min(width * 2 + 4, 60)
    wb.save(report_path)
    log(f"结果报告已保存：{report_path}")
    log(f"完成：文件夹{verb} {moved}/未匹配 {failed}/冲突 {conflict}，"
        f"文件{verb} {moved_f}/未匹配 {failed_f}/冲突 {conflict_f}，源缺失 {missing}")
    return moved, moved_f, failed, failed_f, conflict, conflict_f, missing, report_path


# ---------------------------------------------------------------- 功能二：目录清单

def generate_listing(root_dir, out_xlsx=None, compare_excel=None, log=print):
    import openpyxl
    root_dir = os.path.abspath(root_dir)
    rows = []       # (depth, parts, name, type, size_kb)
    max_depth = 0
    for cur, dirs, files in os.walk(root_dir):
        rel = os.path.relpath(cur, root_dir)
        parts = [] if rel == "." else rel.split(os.sep)
        for d in sorted(dirs):
            p = parts + [d]
            rows.append((len(p), p, d, "文件夹", None))
            max_depth = max(max_depth, len(p))
        for fn in sorted(files):
            p = parts + [fn]
            try:
                size_kb = round(os.path.getsize(os.path.join(cur, fn)) / 1024, 1)
            except OSError:
                size_kb = None
            rows.append((len(p), p, fn, "文件", size_kb))
            max_depth = max(max_depth, len(p))

    # 可选：与分类 Excel 比对（编号列 = 原文件夹名）
    mapping = {}
    if compare_excel:
        mapping = load_classification(compare_excel, log=log)
        log(f"比对分类表 {len(mapping)} 条记录")

    if not out_xlsx:
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        out_xlsx = os.path.join(root_dir, f"目录清单_{ts}.xlsx")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "目录清单"
    header = ["序号", "名称", "类型", "相对路径", "深度"]
    header += [f"第{i}级" for i in range(1, max_depth + 1)]
    header.append("大小(KB)")
    if mapping:
        header += ["是否在分类表中", "分类目标"]
    ws.append(header)

    top_names = set()
    for i, (depth, parts, name, typ, size_kb) in enumerate(rows, start=1):
        rel = os.path.join(*parts)
        row = [i, name, typ, rel, depth]
        row += [parts[j] if j < len(parts) else "" for j in range(max_depth)]
        row.append(size_kb if size_kb is not None else "")
        if mapping:
            if depth == 1 and typ == "文件夹":
                top_names.add(name)
                hit = name if name in mapping else None
                row.append("是" if hit else "否")
                row.append(deepest_path(mapping[hit]) if hit else "")
            else:
                row += ["", ""]
        ws.append(row)

    for col_cells in ws.columns:
        width = max(len(str(c.value)) if c.value is not None else 0 for c in col_cells)
        ws.column_dimensions[col_cells[0].column_letter].width = min(width * 2 + 4, 60)
    wb.save(out_xlsx)
    n_dir = sum(1 for r in rows if r[3] == "文件夹")
    n_file = len(rows) - n_dir
    log(f"清单已保存：{out_xlsx}（文件夹 {n_dir} 个，文件 {n_file} 个）")
    if mapping:
        miss = sorted(n for n in top_names if n not in mapping)
        log(f"顶层文件夹 {len(top_names)} 个，其中 {len(miss)} 个不在分类表中")
    return out_xlsx


# ---------------------------------------------------------------- 功能四：按 Excel 批量重命名

def load_rename_map(excel_path, log=print):
    """读取重命名 Excel，返回 {旧名称: 新名称}。
    自动识别表头（含"旧/原"的列=旧名称，含"新"的列=新名称）；
    识别不到表头时按 A 列=旧名称、B 列=新名称（无表头）处理。"""
    import openpyxl
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    last_err = None
    for ws in wb.worksheets:
        header_row = old_col = new_col = None
        for r_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=10, values_only=True), start=1):
            oc = nc = None
            for c_idx, v in enumerate(row):
                cn = _norm(v)
                if not cn:
                    continue
                if oc is None and any(k in cn for k in ("旧", "原", "old", "before")):
                    oc = c_idx
                    continue
                if nc is None and any(k in cn for k in ("新", "new", "after")):
                    nc = c_idx
            if oc is not None and nc is not None:
                header_row, old_col, new_col = r_idx, oc, nc
                break
        if header_row is None:
            header_row, old_col, new_col = 0, 0, 1
            log(f"Sheet「{ws.title}」未识别到表头，按 A列=旧名称、B列=新名称（无表头）处理")
        else:
            header = [ws.cell(row=header_row, column=c + 1).value for c in range(ws.max_column)]
            log(f"使用 Sheet「{ws.title}」，表头在第 {header_row} 行："
                f"旧名称列=「{header[old_col]}」，新名称列=「{header[new_col]}」")
        mapping = {}
        for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
            old = _cell_to_name(row[old_col] if old_col < len(row) else None)
            new = _cell_to_name(row[new_col] if new_col < len(row) else None)
            if not old or not new:
                continue
            mapping[old] = new
        if mapping:
            return mapping
        last_err = "Sheet 中没有有效的重命名记录"
    raise ValueError(str(last_err) if last_err else "Excel 中没有可用的 Sheet")


def rename_by_excel(excel_path, work_dir, mode="dirs", action="move", out_dir=None, log=print):
    """按 Excel（旧名称->新名称）批量重命名。
    mode: 'dirs' 只重命名子文件夹；'files' 只重命名文件（不限格式）；'both' 两者
    action: 'move' 原位重命名；'copy' 复制到新目录并用新名称（原位置保留）
    文件按"去后缀名"匹配旧名称；新名称不带扩展名时保留原扩展名。
    """
    import openpyxl
    if action == "copy":
        if not out_dir:
            raise ValueError("复制模式需要指定输出目录")
        os.makedirs(out_dir, exist_ok=True)
        base_dir = out_dir
    else:
        action = "move"
        base_dir = work_dir

    mapping = load_rename_map(excel_path, log=log)
    log(f"重命名表共 {len(mapping)} 条记录")
    lower_index = {k.lower(): k for k in mapping}

    do_dirs = mode in ("both", "dirs")
    do_files = mode in ("both", "files")
    mode_text = {"both": "文件夹+文件", "dirs": "仅文件夹", "files": "仅文件"}.get(mode, mode)
    action_text = "复制到新目录" if action == "copy" else "移动（原位重命名）"
    verb = "复制" if action == "copy" else "重命名"
    log(f"整理对象：{mode_text}；处理方式：{action_text}")

    records = []  # (原名称, 新名称, 处理结果, 说明)
    ok = unmatched = conflict = 0

    def _lookup(name):
        if name in mapping:
            return name
        return lower_index.get(name.lower())

    if do_dirs:
        dirs = sorted(d for d in os.listdir(work_dir) if os.path.isdir(os.path.join(work_dir, d)))
        log(f"待处理子文件夹 {len(dirs)} 个")
        for name in dirs:
            key = _lookup(name)
            if key is None:
                unmatched += 1
                records.append((name, "", "未匹配-未处理", "重命名表中无此名称"))
                log(f"[未匹配] {name}")
                continue
            new_name = mapping[key]
            src = os.path.join(work_dir, name)
            dst = os.path.join(base_dir, new_name)
            if os.path.exists(dst):
                conflict += 1
                records.append((name, new_name, f"目标冲突-未{verb}", "目标位置已存在同名文件夹"))
                log(f"[冲突] {name} -> {new_name}：目标已存在")
                continue
            if action == "copy":
                shutil.copytree(src, dst)
            else:
                shutil.move(src, dst)
            ok += 1
            records.append((name, new_name, f"已{verb}", ""))
            log(f"[{verb}] {name} -> {new_name}")

    if do_files:
        files = sorted(f for f in os.listdir(work_dir) if os.path.isfile(os.path.join(work_dir, f)))
        log(f"待处理文件 {len(files)} 个")
        for fn in files:
            stem, ext = os.path.splitext(fn)
            key = _lookup(stem) or _lookup(fn)  # 先按去后缀名，再按完整文件名
            if key is None:
                unmatched += 1
                records.append((fn, "", "未匹配-未处理", "重命名表中无此名称"))
                log(f"[未匹配] {fn}")
                continue
            new = mapping[key]
            # 新名称不带扩展名时保留原扩展名；带扩展名则以新名称为准
            new_name = new if os.path.splitext(new)[1] else new + ext
            src = os.path.join(work_dir, fn)
            dst = os.path.join(base_dir, new_name)
            if os.path.exists(dst):
                conflict += 1
                records.append((fn, new_name, f"目标冲突-未{verb}", "目标位置已存在同名文件"))
                log(f"[冲突] {fn} -> {new_name}：目标已存在")
                continue
            if action == "copy":
                shutil.copy2(src, dst)
            else:
                shutil.move(src, dst)
            ok += 1
            records.append((fn, new_name, f"已{verb}", ""))
            log(f"[{verb}] {fn} -> {new_name}")

    # 结果报告
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    report_path = os.path.join(base_dir, f"重命名结果报告_{ts}.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "重命名结果"
    ws.append(["原名称", "新名称", "处理结果", "说明"])
    for rec in records:
        ws.append(list(rec))
    summary = wb.create_sheet("汇总")
    summary.append(["项目", "数量"])
    summary.append(["整理对象", mode_text])
    summary.append(["处理方式", action_text])
    summary.append([f"成功{verb}", ok])
    summary.append(["未匹配（未处理）", unmatched])
    summary.append(["目标冲突", conflict])
    summary.append(["报告生成时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")])
    for sheet in (ws, summary):
        for col_cells in sheet.columns:
            width = max(len(str(c.value)) if c.value is not None else 0 for c in col_cells)
            sheet.column_dimensions[col_cells[0].column_letter].width = min(width * 2 + 4, 60)
    wb.save(report_path)
    log(f"结果报告已保存：{report_path}")
    log(f"完成：{verb} {ok}，未匹配 {unmatched}，冲突 {conflict}")
    return ok, unmatched, conflict, report_path


# ---------------------------------------------------------------- 功能三：同名文件归组

def group_files(folder, mode="full", prefix_len=5, exts=None, log=print):
    """把顶层文件按文件名分组移入同名文件夹。
    mode='full'   按完整文件名（去扩展名）
    mode='prefix' 按文件名前 prefix_len 个字符
    exts: None 表示全部文件；否则为 {'.jpg', '.shp'} 等小写扩展名集合
    """
    if exts:
        exts = {e.lower() if e.startswith(".") else "." + e.lower() for e in exts}
    files = sorted(f for f in os.listdir(folder) if os.path.isfile(os.path.join(folder, f)))
    moved = skipped = conflict = 0
    for fn in files:
        stem, ext = os.path.splitext(fn)
        if exts and ext.lower() not in exts:
            continue
        key = stem if mode == "full" else stem[:prefix_len]
        key = key.strip()
        if not key:
            skipped += 1
            log(f"[跳过] {fn}：无法确定分组名")
            continue
        dst_dir = os.path.join(folder, key)
        dst = os.path.join(dst_dir, fn)
        if os.path.exists(dst):
            conflict += 1
            log(f"[冲突] {fn}：{key} 中已存在同名文件，未移动")
            continue
        os.makedirs(dst_dir, exist_ok=True)
        shutil.move(os.path.join(folder, fn), dst)
        moved += 1
        log(f"[移动] {fn} -> {key}")
    log(f"完成：移动 {moved} 个文件，跳过 {skipped}，冲突 {conflict}")
    return moved, skipped, conflict


# ---------------------------------------------------------------- GUI

def run_gui():
    import tkinter as tk
    from tkinter import ttk, filedialog, messagebox, scrolledtext
    try:
        from tkinterdnd2 import TkinterDnD, DND_FILES
        RootClass = TkinterDnD.Tk
        dnd_ok = True
    except Exception:
        RootClass = tk.Tk
        dnd_ok = False

    cfg = load_config()
    root = RootClass()
    root.title(APP_NAME)
    try:
        root.iconbitmap(default=_resource_path(os.path.join("assets", "icon.ico")))
    except Exception:
        pass
    root.geometry("860x640")

    nb = ttk.Notebook(root)
    nb.pack(fill=tk.BOTH, expand=True, padx=6, pady=6)

    def parse_drop(data):
        # 拖拽数据可能是 {带空格路径} 或多条；取第一条
        try:
            items = root.tk.splitlist(data)
            return items[0] if items else ""
        except Exception:
            return data.strip("{}")

    def make_path_row(parent, row, label, var, is_file, cfg_key):
        ttk.Label(parent, text=label).grid(row=row, column=0, sticky="w", pady=3)
        ent = ttk.Entry(parent, textvariable=var, width=78)
        ent.grid(row=row, column=1, sticky="ew", padx=4)
        if cfg.get(cfg_key):
            var.set(cfg[cfg_key])

        def browse():
            if is_file:
                p = filedialog.askopenfilename(title="选择文件",
                                               filetypes=[("Excel 文件", "*.xlsx *.xlsm *.xls")])
            else:
                p = filedialog.askdirectory(title="选择文件夹")
            if p:
                var.set(p)
        ttk.Button(parent, text="浏览…", command=browse).grid(row=row, column=2, padx=2)
        if dnd_ok:
            ent.drop_target_register("DND_Files")
            ent.dnd_bind("<<Drop>>", lambda e, v=var: v.set(parse_drop(e.data)))
        return ent

    def make_log(parent, row):
        box = scrolledtext.ScrolledText(parent, height=18, state="disabled")
        box.grid(row=row, column=0, columnspan=3, sticky="nsew", pady=(8, 0))
        parent.rowconfigure(row, weight=1)

        def log(msg):
            def _w():
                box.configure(state="normal")
                box.insert(tk.END, str(msg) + "\n")
                box.see(tk.END)
                box.configure(state="disabled")
            box.after(0, _w)
        return log

    def run_threaded(btn, fn, on_done=None):
        btn.config(state="disabled")

        def worker():
            err = None
            try:
                fn()
            except Exception as e:
                err = e
                traceback.print_exc()
            def fin():
                btn.config(state="normal")
                if err is not None:
                    messagebox.showerror("出错", str(err))
                elif on_done:
                    on_done()
            root.after(0, fin)
        threading.Thread(target=worker, daemon=True).start()

    # ---------------- 选项卡一：按 Excel 分类整理 ----------------
    tab1 = ttk.Frame(nb, padding=10)
    nb.add(tab1, text=" 按Excel分类整理 ")
    tab1.columnconfigure(1, weight=1)
    excel_var1 = tk.StringVar()
    dir_var1 = tk.StringVar()
    make_path_row(tab1, 0, "分类 Excel 文件：", excel_var1, True, "classify_excel")
    make_path_row(tab1, 1, "待处理文件夹：", dir_var1, False, "classify_dir")

    # 整理对象：对文件夹调整 / 对里面的文件调整（二选一）
    obj_var = tk.StringVar(value=cfg.get("classify_obj", "dirs"))
    frm_obj = ttk.LabelFrame(tab1, text="整理对象（二选一）", padding=6)
    frm_obj.grid(row=2, column=0, columnspan=3, sticky="ew", pady=4)
    ttk.Radiobutton(frm_obj, text="对文件夹调整（按文件夹名包含编号匹配，移动子文件夹）",
                    variable=obj_var, value="dirs").grid(row=0, column=0, sticky="w")
    ttk.Radiobutton(frm_obj, text="对里面的文件调整（按文件名包含编号匹配，不限格式，移动文件）",
                    variable=obj_var, value="files").grid(row=1, column=0, sticky="w")

    # 处理方式：移动 / 复制到新目录（二选一）
    act_var = tk.StringVar(value=cfg.get("classify_action", "move"))
    frm_act = ttk.LabelFrame(tab1, text="处理方式（二选一）", padding=6)
    frm_act.grid(row=3, column=0, columnspan=3, sticky="ew", pady=4)
    ttk.Radiobutton(frm_act, text="移动（直接在待处理文件夹内整理，原位置不再保留）",
                    variable=act_var, value="move").grid(row=0, column=0, sticky="w")
    ttk.Radiobutton(frm_act, text="复制到新目录（待处理文件夹保持原样，整理结果输出到下面指定的新目录）",
                    variable=act_var, value="copy").grid(row=1, column=0, sticky="w")

    outdir_var = tk.StringVar()
    make_path_row(tab1, 4, "新目录（复制时必填）：", outdir_var, False, "classify_outdir")

    tip = ("编号可为任意字符串，名称中包含编号即匹配（位置不限）；去向按 Excel 级别目录"
           "（一级/二级/三级…级别数不限，取最深非空级）。未匹配项进入「整理失败」")
    ttk.Label(tab1, text=tip, foreground="#666").grid(row=5, column=0, columnspan=3, sticky="w")
    log1 = make_log(tab1, 7)

    def start_classify():
        ep, wd = excel_var1.get().strip(), dir_var1.get().strip()
        if not os.path.isfile(ep):
            messagebox.showerror("错误", "请选择有效的分类 Excel 文件")
            return
        if not os.path.isdir(wd):
            messagebox.showerror("错误", "请选择有效的待处理文件夹")
            return
        obj = obj_var.get()
        act = act_var.get()
        od = outdir_var.get().strip()
        if act == "copy":
            if not od:
                messagebox.showerror("错误", "复制模式请先选择「新目录」")
                return
            if os.path.abspath(od) == os.path.abspath(wd):
                messagebox.showerror("错误", "新目录不能与待处理文件夹相同")
                return
        obj_text = "文件夹" if obj == "dirs" else "里面的文件"
        act_text = "移动（在待处理文件夹内整理）" if act == "move" else f"复制到新目录：{od}"
        if not messagebox.askyesno("确认", f"整理对象：{obj_text}\n处理方式：{act_text}\n"
                                           "未匹配的会进入「整理失败」。是否继续？"):
            return
        cfg.update({"classify_excel": ep, "classify_dir": wd, "classify_obj": obj,
                    "classify_action": act, "classify_outdir": od})
        save_config(cfg)

        def job():
            log1("=" * 60)
            log1(f"Excel：{ep}")
            log1(f"文件夹：{wd}")
            log1(f"整理对象：{obj_text}；处理方式：{act_text}")
            res = classify(ep, wd, mode=obj, action=act, out_dir=od or None, log=log1)
            log1(f"汇总：文件夹 {res[0]}/未匹配 {res[2]}/冲突 {res[4]}；"
                 f"文件 {res[1]}/未匹配 {res[3]}/冲突 {res[5]}；源缺失 {res[6]}")
        run_threaded(btn1, job, on_done=lambda: messagebox.showinfo("完成", "分类整理完成，详见日志与结果报告"))

    btn1 = ttk.Button(tab1, text="开始分类整理", command=start_classify)
    btn1.grid(row=6, column=0, columnspan=3, pady=6)

    # ---------------- 选项卡二：目录清单 ----------------
    tab2 = ttk.Frame(nb, padding=10)
    nb.add(tab2, text=" 批量获取文件夹名 ")
    tab2.columnconfigure(1, weight=1)
    dir_var2 = tk.StringVar()
    cmp_var = tk.StringVar()
    out_var = tk.StringVar()
    make_path_row(tab2, 0, "要扫描的文件夹：", dir_var2, False, "scan_dir")
    make_path_row(tab2, 1, "比对 Excel（可空）：", cmp_var, True, "scan_compare")
    make_path_row(tab2, 2, "清单保存位置（可空）：", out_var, True, "")
    ttk.Label(tab2, text="递归列出全部子文件夹与文件；填比对 Excel 后会标注每个顶层文件夹的归类去向",
              foreground="#666").grid(row=3, column=0, columnspan=3, sticky="w")
    log2 = make_log(tab2, 5)

    def start_scan():
        wd = dir_var2.get().strip()
        if not os.path.isdir(wd):
            messagebox.showerror("错误", "请选择有效的文件夹")
            return
        cmp_x = cmp_var.get().strip() or None
        if cmp_x and not os.path.isfile(cmp_x):
            messagebox.showerror("错误", "比对 Excel 无效")
            return
        out_x = out_var.get().strip() or None
        cfg.update({"scan_dir": wd, "scan_compare": cmp_var.get().strip()})
        save_config(cfg)

        def job():
            log2("=" * 60)
            log2(f"扫描：{wd}")
            generate_listing(wd, out_x, cmp_x, log=log2)
        run_threaded(btn2, job, on_done=lambda: messagebox.showinfo("完成", "目录清单已生成"))

    btn2 = ttk.Button(tab2, text="生成目录清单", command=start_scan)
    btn2.grid(row=4, column=0, columnspan=3, pady=6)

    # ---------------- 选项卡二·五：批量重命名（按 Excel 旧名称->新名称） ----------------
    tab2b = ttk.Frame(nb, padding=10)
    nb.add(tab2b, text=" 批量重命名 ")
    tab2b.columnconfigure(1, weight=1)
    excel_var_r = tk.StringVar()
    dir_var_r = tk.StringVar()
    make_path_row(tab2b, 0, "重命名 Excel 文件：", excel_var_r, True, "rename_excel")
    make_path_row(tab2b, 1, "待处理文件夹：", dir_var_r, False, "rename_dir")

    obj_var_r = tk.StringVar(value=cfg.get("rename_obj", "dirs"))
    frm_obj_r = ttk.LabelFrame(tab2b, text="整理对象（二选一）", padding=6)
    frm_obj_r.grid(row=2, column=0, columnspan=3, sticky="ew", pady=4)
    ttk.Radiobutton(frm_obj_r, text="对文件夹调整（只重命名子文件夹名称）",
                    variable=obj_var_r, value="dirs").grid(row=0, column=0, sticky="w")
    ttk.Radiobutton(frm_obj_r, text="对里面的文件调整（只重命名文件名称，不限格式）",
                    variable=obj_var_r, value="files").grid(row=1, column=0, sticky="w")

    act_var_r = tk.StringVar(value=cfg.get("rename_action", "move"))
    frm_act_r = ttk.LabelFrame(tab2b, text="处理方式（二选一）", padding=6)
    frm_act_r.grid(row=3, column=0, columnspan=3, sticky="ew", pady=4)
    ttk.Radiobutton(frm_act_r, text="移动（直接在待处理文件夹内重命名，原名称不再保留）",
                    variable=act_var_r, value="move").grid(row=0, column=0, sticky="w")
    ttk.Radiobutton(frm_act_r, text="复制到新目录（待处理文件夹保持原样，用新名称复制到下面指定的新目录）",
                    variable=act_var_r, value="copy").grid(row=1, column=0, sticky="w")

    outdir_var_r = tk.StringVar()
    make_path_row(tab2b, 4, "新目录（复制时必填）：", outdir_var_r, False, "rename_outdir")

    tip_r = ("Excel 需两列：一列旧名称、一列新名称（表头含“旧/原”“新”字样可自动识别，"
             "无表头则按 A列=旧名称、B列=新名称）。文件按去后缀名匹配，新名称不带扩展名时保留原扩展名")
    ttk.Label(tab2b, text=tip_r, foreground="#666", wraplength=800, justify="left").grid(
        row=5, column=0, columnspan=3, sticky="w")
    log_r = make_log(tab2b, 7)

    def start_rename():
        ep, wd = excel_var_r.get().strip(), dir_var_r.get().strip()
        if not os.path.isfile(ep):
            messagebox.showerror("错误", "请选择有效的重命名 Excel 文件")
            return
        if not os.path.isdir(wd):
            messagebox.showerror("错误", "请选择有效的待处理文件夹")
            return
        obj = obj_var_r.get()
        act = act_var_r.get()
        od = outdir_var_r.get().strip()
        if act == "copy":
            if not od:
                messagebox.showerror("错误", "复制模式请先选择「新目录」")
                return
            if os.path.abspath(od) == os.path.abspath(wd):
                messagebox.showerror("错误", "新目录不能与待处理文件夹相同")
                return
        obj_text = "文件夹" if obj == "dirs" else "里面的文件"
        act_text = "移动（原位重命名）" if act == "move" else f"复制到新目录：{od}"
        if not messagebox.askyesno("确认", f"整理对象：{obj_text}\n处理方式：{act_text}\n是否继续？"):
            return
        cfg.update({"rename_excel": ep, "rename_dir": wd, "rename_obj": obj,
                    "rename_action": act, "rename_outdir": od})
        save_config(cfg)

        def job():
            log_r("=" * 60)
            log_r(f"Excel：{ep}")
            log_r(f"文件夹：{wd}")
            log_r(f"整理对象：{obj_text}；处理方式：{act_text}")
            res = rename_by_excel(ep, wd, mode=obj, action=act, out_dir=od or None, log=log_r)
            log_r(f"汇总：成功 {res[0]}，未匹配 {res[1]}，冲突 {res[2]}")
        run_threaded(btn_r, job, on_done=lambda: messagebox.showinfo("完成", "批量重命名完成，详见日志与结果报告"))

    btn_r = ttk.Button(tab2b, text="开始批量重命名", command=start_rename)
    btn_r.grid(row=6, column=0, columnspan=3, pady=6)

    # ---------------- 选项卡三：同名文件归组 ----------------
    tab3 = ttk.Frame(nb, padding=10)
    nb.add(tab3, text=" 同名文件归组 ")
    tab3.columnconfigure(1, weight=1)
    dir_var3 = tk.StringVar()
    make_path_row(tab3, 0, "文件所在文件夹：", dir_var3, False, "group_dir")

    mode_var = tk.StringVar(value=cfg.get("group_mode", "full"))
    frm_mode = ttk.LabelFrame(tab3, text="分组方式", padding=6)
    frm_mode.grid(row=1, column=0, columnspan=3, sticky="ew", pady=4)
    ttk.Radiobutton(frm_mode, text="按完整文件名（同名不同扩展名归入同一文件夹）",
                    variable=mode_var, value="full").grid(row=0, column=0, sticky="w")
    ttk.Radiobutton(frm_mode, text="按文件名前 N 位", variable=mode_var,
                    value="prefix").grid(row=1, column=0, sticky="w")
    prefix_var = tk.StringVar(value=str(cfg.get("group_prefix", 5)))
    ttk.Label(frm_mode, text="N =").grid(row=1, column=1, sticky="e")
    ttk.Entry(frm_mode, textvariable=prefix_var, width=6).grid(row=1, column=2, sticky="w")

    ext_var = tk.StringVar(value=cfg.get("group_exts", ""))
    ttk.Label(tab3, text="扩展名过滤：").grid(row=2, column=0, sticky="w", pady=3)
    ttk.Entry(tab3, textvariable=ext_var, width=40).grid(row=2, column=1, sticky="w", padx=4)
    ttk.Label(tab3, text="如 jpg,shp；留空=全部文件", foreground="#666").grid(row=2, column=2, sticky="w")
    log3 = make_log(tab3, 4)

    def start_group():
        wd = dir_var3.get().strip()
        if not os.path.isdir(wd):
            messagebox.showerror("错误", "请选择有效的文件夹")
            return
        try:
            n = max(1, int(prefix_var.get()))
        except ValueError:
            messagebox.showerror("错误", "N 必须是正整数")
            return
        exts = [e.strip() for e in re.split(r"[,，;；\s]+", ext_var.get()) if e.strip()] or None
        mode = mode_var.get()
        if not messagebox.askyesno("确认", f"即将把 {wd} 中的文件按文件名分组移入子文件夹，是否继续？"):
            return
        cfg.update({"group_dir": wd, "group_mode": mode, "group_prefix": n,
                    "group_exts": ext_var.get()})
        save_config(cfg)

        def job():
            log3("=" * 60)
            log3(f"文件夹：{wd}，方式：{'完整文件名' if mode == 'full' else f'前 {n} 位'}，扩展名：{exts or '全部'}")
            group_files(wd, mode=mode, prefix_len=n, exts=exts, log=log3)
        run_threaded(btn3, job, on_done=lambda: messagebox.showinfo("完成", "文件归组完成"))

    btn3 = ttk.Button(tab3, text="开始归组", command=start_group)
    btn3.grid(row=3, column=0, columnspan=3, pady=6)

    if dnd_ok:
        root.statusbar = ttk.Label(root, text="提示：可直接把文件/文件夹拖入输入框", anchor="w", foreground="#666")
    else:
        root.statusbar = ttk.Label(root, text="提示：拖拽组件不可用，请使用浏览按钮", anchor="w", foreground="#666")
    root.statusbar.pack(side=tk.BOTTOM, fill=tk.X)

    root.mainloop()


# ---------------------------------------------------------------- CLI

def _usage():
    print("""用法：
  文件夹整理工具.exe                                   打开图形界面
  文件夹整理工具.exe classify <分类Excel> <待处理文件夹> [dirs|files|both] [move|copy] [输出目录]
  文件夹整理工具.exe rename <重命名Excel> <待处理文件夹> [dirs|files|both] [move|copy] [输出目录]
  文件夹整理工具.exe scan <文件夹> [清单输出.xlsx] [比对Excel]
  文件夹整理工具.exe group <文件夹> [full|prefix] [N] [扩展名如 jpg,shp]""")


if __name__ == "__main__":
    args = sys.argv[1:]
    if not args:
        run_gui()
    elif args[0] == "classify" and len(args) >= 3:
        classify(args[1], args[2],
                 mode=args[3] if len(args) > 3 else "both",
                 action=args[4] if len(args) > 4 else "move",
                 out_dir=args[5] if len(args) > 5 else None)
    elif args[0] == "rename" and len(args) >= 3:
        rename_by_excel(args[1], args[2],
                        mode=args[3] if len(args) > 3 else "dirs",
                        action=args[4] if len(args) > 4 else "move",
                        out_dir=args[5] if len(args) > 5 else None)
    elif args[0] == "scan" and len(args) >= 2:
        generate_listing(args[1],
                         args[2] if len(args) > 2 else None,
                         args[3] if len(args) > 3 else None)
    elif args[0] == "group" and len(args) >= 2:
        mode = args[2] if len(args) > 2 else "full"
        n = int(args[3]) if len(args) > 3 else 5
        exts = [e.strip() for e in re.split(r"[,，;；\s]+", args[4]) if e.strip()] if len(args) > 4 else None
        group_files(args[1], mode=mode, prefix_len=n, exts=exts)
    else:
        _usage()
